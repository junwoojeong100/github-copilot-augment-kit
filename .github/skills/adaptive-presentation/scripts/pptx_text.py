"""Read semantic visible text without applying language or geometry policies."""

from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def _chart_flag(element, name: str, default: bool = False) -> bool:
    flag = element.find(qn(f"c:{name}")) if element is not None else None
    return flag.get("val", "1") in {"1", "true"} if flag is not None else default


def _chart_text_values(element):
    if element is None:
        return
    for paragraph in element.findall(f".//{qn('a:p')}"):
        text = "".join(
            value.text or "" for value in paragraph.findall(f".//{qn('a:t')}")
        )
        if text.strip():
            yield text
    for value in element.findall(f".//{qn('c:v')}"):
        if value.text and value.text.strip():
            yield value.text


def _chart_texts(chart):
    chart_xml = chart._chartSpace.find(qn("c:chart"))
    yield from _chart_text_values(chart_xml.find(qn("c:title")))
    plot_area = chart_xml.find(qn("c:plotArea"))
    visible_category_axes: set[str] = set()
    visible_series_axes: set[str] = set()
    axis_tags = {qn(f"c:{name}") for name in ("catAx", "dateAx", "valAx", "serAx")}
    for axis in plot_area:
        if axis.tag not in axis_tags or _chart_flag(axis, "delete"):
            continue
        yield from _chart_text_values(axis.find(qn("c:title")))
        position = axis.find(qn("c:tickLblPos"))
        axis_id = axis.find(qn("c:axId"))
        if axis_id is not None and (position is None or position.get("val") != "none"):
            if axis.tag in {qn("c:catAx"), qn("c:dateAx")}:
                visible_category_axes.add(axis_id.attrib["val"])
            elif axis.tag == qn("c:serAx"):
                visible_series_axes.add(axis_id.attrib["val"])
    legend = chart_xml.find(qn("c:legend"))
    deleted_legend_entries = {
        entry.find(qn("c:idx")).attrib["val"]
        for entry in legend.findall(qn("c:legendEntry"))
        if _chart_flag(entry, "delete")
    } if legend is not None else set()
    for plot in plot_area:
        axes = {axis.attrib["val"] for axis in plot.findall(qn("c:axId"))}
        category_legend = plot.tag in {
            qn("c:pieChart"), qn("c:pie3DChart"),
            qn("c:doughnutChart"), qn("c:ofPieChart"),
        }
        plot_labels = plot.find(qn("c:dLbls"))
        for series in plot.findall(qn("c:ser")):
            name = " ".join(_chart_text_values(series.find(qn("c:tx"))))
            categories: dict[str, list[str]] = {}
            category = series.find(qn("c:cat"))
            if category is not None:
                for point in category.findall(f".//{qn('c:pt')}"):
                    categories.setdefault(point.attrib["idx"], []).extend(
                        _chart_text_values(point)
                    )
            if axes & visible_category_axes:
                for texts in categories.values():
                    yield from texts
            if axes & visible_series_axes:
                yield name
            if legend is not None:
                if category_legend:
                    for index, texts in categories.items():
                        if index not in deleted_legend_entries:
                            yield from texts
                elif series.find(qn("c:idx")).attrib["val"] not in deleted_legend_entries:
                    yield name
            series_labels = series.find(qn("c:dLbls"))
            if plot_labels is None and series_labels is None:
                continue
            point_labels = {}
            for labels in (plot_labels, series_labels):
                if labels is not None:
                    point_labels.update({
                        label.find(qn("c:idx")).attrib["val"]: label
                        for label in labels.findall(qn("c:dLbl"))
                    })
            point_indices = set(categories) | set(point_labels)
            for data_name in ("val", "xVal", "yVal"):
                data = series.find(qn(f"c:{data_name}"))
                if data is not None:
                    point_indices.update(
                        point.attrib["idx"] for point in data.findall(f".//{qn('c:pt')}")
                    )
            for index in sorted(point_indices):
                label = point_labels.get(index)
                if _chart_flag(label, "delete"):
                    continue
                text = label.find(qn("c:tx")) if label is not None else None
                if text is not None:
                    yield from _chart_text_values(text)
                    continue
                for flag, texts in (
                    ("showCatName", categories.get(index, [])),
                    ("showSerName", [name]),
                ):
                    default = _chart_flag(
                        series_labels, flag, _chart_flag(plot_labels, flag)
                    )
                    if _chart_flag(label, flag, default):
                        yield from texts


def _shape_texts(shape, *, inherited_top: int | None = None):
    top = inherited_top if inherited_top is not None else getattr(shape, "top", None)
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            yield from _shape_texts(child, inherited_top=top)
    elif getattr(shape, "has_chart", False):
        for text in _chart_texts(shape.chart):
            if text.strip():
                yield top, text
    elif getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    yield top, cell.text
    elif getattr(shape, "has_text_frame", False) and shape.text.strip():
        yield top, shape.text


def iter_slide_texts(slide):
    """Yield ``(top_emu, text)`` for displayed shape, table, and chart text."""
    for shape in slide.shapes:
        yield from _shape_texts(shape)
