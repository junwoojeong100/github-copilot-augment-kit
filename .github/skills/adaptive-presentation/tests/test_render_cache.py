from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch


SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

try:
    import pymupdf as fitz
except ImportError:
    import fitz
import render_cache  # noqa: E402
import render_pptx  # noqa: E402
import verify_deck  # noqa: E402
import visual_review  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402


class RenderCacheTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name).resolve()
        self.deck = self.root / "deck.pptx"
        self.deck.write_bytes(b"deck")
        self.qa = self.root / "qa"
        render_pptx.claim_output_dir(self.qa)
        self.pdf = self.qa / "deck.pdf"
        self.pdf.write_bytes(b"pdf")
        self.sheet = self.qa / "contact-01.jpg"
        self.sheet.write_bytes(b"preview")
        self.args = verify_deck.render_namespace(self.deck, self.qa)
        self.current = {"deck": "a" * 64, "environment": "b" * 64, "options": "c" * 64}
        self.manifest = {
            "deck": str(self.deck),
            "deck_sha256": render_pptx.sha256_file(self.deck),
            "pdf": str(self.pdf),
            "pdf_sha256": render_pptx.sha256_file(self.pdf),
            "total_slides": 1,
            "rendered_slides": [1],
            "contact_sheets": [str(self.sheet)],
            "slide_images": [],
            "kept_pdf": True,
        }
        self.manifest_path = self.qa / "manifest.json"
        self.manifest_path.write_text(json.dumps(self.manifest), encoding="utf-8")
        render_cache.store(self.args, self.current, self.manifest)

    def test_matching_artifacts_are_reused_without_mutation(self):
        original = self.manifest_path.read_bytes()
        manifest, reason = render_cache.load(self.args, self.current)
        self.assertEqual(manifest, self.manifest)
        self.assertIn("matching", reason)
        self.assertEqual(self.manifest_path.read_bytes(), original)

    def test_changed_input_environment_and_options_are_explicit_misses(self):
        for key in self.current:
            changed = {**self.current, key: "d" * 64}
            with self.subTest(key=key):
                manifest, reason = render_cache.load(self.args, changed)
                self.assertIsNone(manifest)
                self.assertEqual(reason, f"changed {key}")

    def test_signature_tracks_deck_path_bytes_and_render_options(self):
        with patch.object(render_cache, "environment_fingerprint", return_value="b" * 64):
            with patch.object(render_pptx, "find_soffice", return_value="soffice"):
                original = render_cache.signature(self.args)
                self.deck.write_bytes(b"changed")
                self.assertNotEqual(original["deck"], render_cache.signature(self.args)["deck"])
                self.deck.write_bytes(b"deck")
                other = self.root / "copy.pptx"
                other.write_bytes(b"deck")
                self.args.deck = other
                self.assertNotEqual(original["deck"], render_cache.signature(self.args)["deck"])
                self.args.deck = self.deck
                self.args.scale = 2.0
                self.assertNotEqual(original["options"], render_cache.signature(self.args)["options"])

    def test_missing_cache_is_a_miss_but_corrupt_metadata_is_an_error(self):
        metadata = self.qa / render_cache.CACHE_NAME
        metadata.unlink()
        self.assertEqual(render_cache.load(self.args, self.current), (None, "no recorded render"))
        for content in ("{", "[]", '{"schemaVersion": 1, "signature": {}}'):
            metadata.write_text(content, encoding="utf-8")
            with self.subTest(content=content):
                with self.assertRaises(RuntimeError):
                    render_cache.load(self.args, self.current)

    def test_changed_or_missing_artifacts_are_not_silent_cache_misses(self):
        for artifact in (self.pdf, self.sheet, self.manifest_path):
            original = artifact.read_bytes()
            with self.subTest(artifact=artifact.name):
                artifact.write_bytes(b"tampered")
                with self.assertRaises(RuntimeError):
                    render_cache.load(self.args, self.current)
                artifact.unlink()
                with self.assertRaises(RuntimeError):
                    render_cache.load(self.args, self.current)
                artifact.write_bytes(original)

    def test_cache_rejects_unowned_directories_and_linked_files(self):
        marker = self.qa / render_pptx.OUTPUT_MARKER
        marker.unlink()
        with self.assertRaisesRegex(RuntimeError, "owned"):
            render_cache.load(self.args, self.current)
        marker.write_text(render_pptx.OUTPUT_MARKER_CONTENT, encoding="utf-8")
        original = self.sheet.read_bytes()
        outside = self.root / "outside.jpg"
        outside.write_bytes(original)
        self.sheet.unlink()
        self.sheet.symlink_to(outside)
        with self.assertRaises(RuntimeError):
            render_cache.load(self.args, self.current)
        self.sheet.unlink()
        self.sheet.hardlink_to(outside)
        with self.assertRaisesRegex(RuntimeError, "linked"):
            render_cache.load(self.args, self.current)
        self.assertEqual(outside.read_bytes(), original)

    def test_partial_renders_and_external_preview_paths_cannot_be_recorded(self):
        for changes in (
            {"total_slides": 2},
            {"rendered_slides": [True]},
            {"contact_sheets": []},
            {"contact_sheets": [str(self.root / "outside.jpg")]},
            {"contact_sheets": [str(self.pdf)]},
        ):
            manifest = {**self.manifest, **changes}
            self.manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            with self.subTest(changes=changes):
                with self.assertRaises(RuntimeError):
                    render_cache.store(self.args, self.current, manifest)

    def test_resource_inventory_follows_links_without_cycles(self):
        fonts = self.root / "fonts"
        fonts.mkdir()
        font = fonts / "font.ttf"
        font.write_bytes(b"font")
        (fonts / "loop").symlink_to(fonts, target_is_directory=True)
        self.assertEqual(render_cache._resource_files([fonts]), {font})

    def test_environment_tracks_font_contents_tool_version_and_locale(self):
        executable = self.root / "soffice"
        executable.write_bytes(b"executable")
        font = self.root / "font.ttf"
        font.write_bytes(b"font")
        with (
            patch.object(render_cache, "_resource_files", return_value={font}),
            patch.object(render_cache, "_native_font_inventory", return_value=({font}, ["Font"])) as native,
            patch.object(render_cache.shutil, "which", return_value=None),
            patch.object(render_cache, "_run_probe", return_value="LibreOffice 1") as probe,
        ):
            original = render_cache.environment_fingerprint(str(executable))
            font.write_bytes(b"changed font")
            self.assertNotEqual(original, render_cache.environment_fingerprint(str(executable)))
            font.write_bytes(b"font")
            probe.return_value = "LibreOffice 2"
            self.assertNotEqual(original, render_cache.environment_fingerprint(str(executable)))
            probe.return_value = "LibreOffice 1"
            native.return_value = ({font}, ["Different registered font"])
            self.assertNotEqual(original, render_cache.environment_fingerprint(str(executable)))
            native.return_value = ({font}, ["Font"])
            with patch.dict(os.environ, {"LC_ALL": "different-locale"}):
                self.assertNotEqual(original, render_cache.environment_fingerprint(str(executable)))

    def test_windows_registry_tracks_external_fonts_and_substitution_state(self):
        font = self.root / "external.ttf"
        font.write_bytes(b"font")
        records = {
            "Fonts": [("External font", str(font), 1)],
            "FontSubstitutes": [("Arial", "External font", 1)],
            "SystemLink": [("Fallback", b"\x00\x01", 3)],
        }

        class Key:
            def __init__(self, values):
                self.values = values

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return None

        def open_key(hive, path):
            if hive == "user":
                raise FileNotFoundError(path)
            return Key(records[path.rsplit("\\", 1)[-1]])

        registry = SimpleNamespace(
            HKEY_LOCAL_MACHINE="machine", HKEY_CURRENT_USER="user",
            OpenKey=open_key, QueryInfoKey=lambda key: (0, len(key.values), 0),
            EnumValue=lambda key, index: key.values[index],
        )
        with (
            patch.object(render_cache.sys, "platform", "win32"),
            patch.dict(sys.modules, {"winreg": registry}),
        ):
            files, state = render_cache._native_font_inventory()
            self.assertEqual(files, {font})
            self.assertEqual(len(state), 3)
            records["FontSubstitutes"] = [("Arial", "Different font", 1)]
            self.assertNotEqual(state, render_cache._native_font_inventory()[1])

    def test_unavailable_coretext_inventory_is_an_explicit_error(self):
        text, foundation = Mock(), Mock()
        text.CTFontManagerCopyAvailableFontURLs.return_value = None
        with (
            patch.object(render_cache.sys, "platform", "darwin"),
            patch.object(render_cache.ctypes, "CDLL", side_effect=[text, foundation]),
        ):
            with self.assertRaisesRegex(RuntimeError, "CoreText font inventory"):
                render_cache._native_font_inventory()


class RenderReuseVerificationTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name).resolve()
        self.deck = self.root / "deck.pptx"
        self.prs = Presentation()
        self.prs.slide_width = Inches(10)
        self.prs.slide_height = Inches(7.5)
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        for text, top, size in (
            ("Decision title", 0.7, 30),
            ("This body sentence provides evidence for the decision.", 2, 15),
        ):
            shape = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(8), Inches(1))
            run = shape.text_frame.paragraphs[0].add_run()
            run.text = text
            run.font.name = "Arial"
            run.font.size = Pt(size)
        self.prs.save(self.deck)
        self.args = verify_deck.build_parser().parse_args([
            str(self.deck), "--out", str(self.root / "verify"), "--reuse-render",
        ])
        self.environment = self.enterContext(patch.object(
            render_cache, "environment_fingerprint", return_value="a" * 64,
        ))
        self.converter = self.enterContext(patch.object(
            render_pptx, "convert_to_pdf", side_effect=self.make_pdf,
        ))
        self.enterContext(patch.object(render_pptx, "find_soffice", return_value="soffice"))

    def make_pdf(self, deck, out, soffice, timeout):
        path = out / f"{deck.stem}.pdf"
        with fitz.open() as document:
            page = document.new_page(width=720, height=540)
            page.insert_text((58, 75), "Decision title", fontsize=30)
            page.insert_text(
                (58, 160), "This body sentence provides evidence for the decision.",
                fontsize=15,
            )
            document.save(path)
        return path

    def test_full_conversion_runs_once_but_audit_and_review_are_rechecked(self):
        with patch.object(
            verify_deck.audit_pptx, "audit", wraps=verify_deck.audit_pptx.audit,
        ) as audit:
            first = verify_deck.verify(self.args)
            self.assertTrue(first["passed"], first["audit_failures"])
            self.assertFalse(first["render_cache"]["reused"])
            self.args.require_visual_review = True
            second = verify_deck.verify(self.args)
            self.assertTrue(second["render_cache"]["reused"])
            self.assertFalse(second["passed"])
            self.assertIn("required", second["visual_review_failure"])
            self.assertEqual(self.converter.call_count, 1)
            self.assertEqual(audit.call_count, 2)
            self.assertEqual(
                first["full_render"]["pdf_sha256"], second["full_render"]["pdf_sha256"],
            )
            evidence = self.root / "visual-review.json"
            evidence.write_text(json.dumps({
                "schemaVersion": 1,
                "deckSha256": render_pptx.sha256_file(self.deck),
                "renderCacheSha256": second["render_cache"]["sha256"],
                "reviewer": "Synthetic test fixture",
                "reviewedSlides": "all",
                "notes": "Synthetic evidence for verifier regression testing.",
            }), encoding="utf-8")
            self.args.visual_review = evidence
            third = verify_deck.verify(self.args)
            self.assertTrue(third["passed"], third["audit_failures"])
            self.assertTrue(third["render_cache"]["reused"])
            self.assertEqual(self.converter.call_count, 1)
            self.assertEqual(audit.call_count, 3)

    def test_environment_change_invalidates_old_visual_approval(self):
        first = verify_deck.verify(self.args)
        evidence = self.root / "visual-review.json"
        evidence.write_text(json.dumps({
            "schemaVersion": 1,
            "deckSha256": render_pptx.sha256_file(self.deck),
            "renderCacheSha256": first["render_cache"]["sha256"],
            "reviewer": "Synthetic test fixture",
            "reviewedSlides": "all",
            "notes": "Synthetic evidence for the original render environment.",
        }), encoding="utf-8")
        self.args.require_visual_review = True
        self.args.visual_review = evidence
        self.environment.return_value = "b" * 64
        second = verify_deck.verify(self.args)
        self.assertFalse(second["passed"])
        self.assertFalse(second["render_cache"]["reused"])
        self.assertIn("render cache", second["visual_review_failure"])
        self.assertNotEqual(first["render_cache"]["sha256"], second["render_cache"]["sha256"])

    def test_revision_specific_reviews_preserve_old_evidence_and_verify_the_new_revision(self):
        verify_deck.verify(self.args)
        cache = self.args.out / "qa" / render_cache.CACHE_NAME
        first_evidence = self.root / "visual-review-r001.json"
        visual_review.write_visual_review(
            self.deck, first_evidence,
            reviewer="Synthetic test fixture",
            notes="Synthetic review of the first rendered revision.",
            render_cache=cache,
        )
        original_evidence = first_evidence.read_bytes()
        self.args.require_visual_review = True
        self.args.visual_review = first_evidence
        self.assertTrue(verify_deck.verify(self.args)["passed"])

        self.prs.slides[0].notes_slide.notes_text_frame.text = "Revised speaker notes."
        self.prs.save(self.deck)
        stale = verify_deck.verify(self.args)
        self.assertFalse(stale["passed"])
        self.assertIn("current PPTX revision", stale["visual_review_failure"])
        with self.assertRaises(FileExistsError):
            visual_review.write_visual_review(
                self.deck, first_evidence,
                reviewer="Synthetic test fixture",
                notes="Synthetic review of the revised rendered deck.",
                render_cache=cache,
            )

        second_evidence = self.root / "visual-review-r002.json"
        visual_review.write_visual_review(
            self.deck, second_evidence,
            reviewer="Synthetic test fixture",
            notes="Synthetic review of the revised rendered deck.",
            render_cache=cache,
        )
        self.args.visual_review = second_evidence
        final = verify_deck.verify(self.args)
        self.assertTrue(final["passed"], final["audit_failures"])
        self.assertTrue(final["render_cache"]["reused"])
        self.assertEqual(self.converter.call_count, 2)
        self.assertEqual(first_evidence.read_bytes(), original_evidence)
        self.assertNotEqual(first_evidence.read_bytes(), second_evidence.read_bytes())

    def test_stricter_qa_cannot_reuse_an_old_pass_decision(self):
        first = verify_deck.verify(self.args)
        self.assertTrue(first["passed"], first["audit_failures"])
        self.args.min_body_pt = 24
        self.args.fail_small_text = True
        second = verify_deck.verify(self.args)
        self.assertTrue(second["render_cache"]["reused"])
        self.assertFalse(second["automated_passed"])
        self.assertTrue(second["audit_failures"])
        self.assertEqual(self.converter.call_count, 1)

    def test_changed_deck_and_environment_force_fresh_rendering(self):
        verify_deck.verify(self.args)
        self.prs.slides[0].notes_slide.notes_text_frame.text = "Changed notes."
        self.prs.save(self.deck)
        changed_deck = verify_deck.verify(self.args)
        self.assertFalse(changed_deck["render_cache"]["reused"])
        self.assertIn("deck", changed_deck["render_cache"]["reason"])
        self.environment.return_value = "b" * 64
        changed_environment = verify_deck.verify(self.args)
        self.assertFalse(changed_environment["render_cache"]["reused"])
        self.assertIn("environment", changed_environment["render_cache"]["reason"])
        self.assertEqual(self.converter.call_count, 3)

    def test_corrupt_cached_pdf_fails_before_any_rerender_or_cleanup(self):
        first = verify_deck.verify(self.args)
        pdf = Path(first["full_render"]["pdf"])
        pdf.write_bytes(b"partial")
        with self.assertRaises(RuntimeError):
            verify_deck.verify(self.args)
        self.assertEqual(pdf.read_bytes(), b"partial")
        self.assertEqual(self.converter.call_count, 1)

    def test_environment_changes_during_render_are_not_cached(self):
        self.environment.side_effect = ["a" * 64, "b" * 64]
        with self.assertRaisesRegex(RuntimeError, "changed during verification"):
            verify_deck.verify(self.args)
        self.assertFalse((self.args.out / "qa" / render_cache.CACHE_NAME).exists())

    def test_cache_metadata_changes_during_reuse_are_not_accepted(self):
        verify_deck.verify(self.args)
        original_audit = verify_deck.audit_pptx.audit

        def mutate_cache(args):
            result = original_audit(args)
            path = self.args.out / "qa" / render_cache.CACHE_NAME
            metadata = json.loads(path.read_text())
            metadata["signature"]["options"] = "f" * 64
            path.write_text(json.dumps(metadata), encoding="utf-8")
            return result

        with patch.object(verify_deck.audit_pptx, "audit", side_effect=mutate_cache):
            with self.assertRaisesRegex(RuntimeError, "cache changed during verification"):
                verify_deck.verify(self.args)
        self.assertEqual(self.converter.call_count, 1)

    def test_default_behavior_still_renders_fresh_without_environment_probes(self):
        self.args.reuse_render = False
        verify_deck.verify(self.args)
        result = verify_deck.verify(self.args)
        self.assertFalse(result["render_cache"]["enabled"])
        self.assertEqual(self.converter.call_count, 2)
        self.environment.assert_not_called()

    def test_reuse_does_not_overwrite_a_linked_input_through_cached_audit(self):
        verify_deck.verify(self.args)
        audit = self.args.out / "qa" / "audit.json"
        audit.unlink()
        original = self.deck.read_bytes()
        audit.hardlink_to(self.deck)
        result = verify_deck.verify(self.args)
        self.assertTrue(result["render_cache"]["reused"])
        self.assertEqual(self.deck.read_bytes(), original)

    def test_report_alias_cannot_modify_a_reused_pdf(self):
        first = verify_deck.verify(self.args)
        pdf = Path(first["full_render"]["pdf"])
        original = pdf.read_bytes()
        report = self.args.out / "verification-report.json"
        report.unlink()
        report.symlink_to(pdf)
        second = verify_deck.verify(self.args)
        self.assertTrue(second["render_cache"]["reused"])
        self.assertFalse(report.is_symlink())
        self.assertEqual(pdf.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
